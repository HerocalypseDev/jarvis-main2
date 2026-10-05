"""PowerPoint (.pptx) files from simple Markdown (2026-09-27, feature batch D3), used by write_file for
.pptx names, the same way jarvis_docx handles .docx. Local only (python-pptx, already installed).

Markdown: each "# Title" starts a slide (the first one is the title slide; a plain line under it is its
subtitle); "- bullet" lines (indent 2 spaces to nest) are the slide's points; plain lines become points
too; a line starting "Notes:" goes into the speaker notes. A line of just --- also starts a new slide.
append=true adds slides to an existing file.
"""

from __future__ import annotations

import re
from pathlib import Path


def parse(markdown: str) -> list[dict]:
    slides: list[dict] = []
    cur = None
    for raw in (markdown or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.strip() == "---":
            cur = None
            continue
        m = re.match(r"^\s*#{1,6}\s+(.*)$", line)
        if m:
            cur = {"title": m.group(1).strip(), "points": [], "notes": []}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": "", "points": [], "notes": []}
            slides.append(cur)
        n = re.match(r"^\s*notes?:\s*(.*)$", line, re.I)
        if n:
            cur["notes"].append(n.group(1))
            continue
        b = re.match(r"^(\s*)[-*•]\s+(.*)$", line)
        if b:
            cur["points"].append((min(len(b.group(1).replace("\t", "  ")) // 2, 4), b.group(2).strip()))
        else:
            cur["points"].append((0, re.sub(r"^\s*\d+[.)]\s+", "", line).strip()))
    return slides


def _plain(text: str) -> str:
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`", lambda m: next(g for g in m.groups() if g), text)


def write(path: Path, markdown: str, append: bool = False) -> int:
    from pptx import Presentation
    markdown = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", markdown or "")  # XML can't hold them
    prs = Presentation(str(path)) if append and Path(path).exists() else Presentation()
    first_new = len(prs.slides) == 0
    for i, s in enumerate(parse(markdown)):
        title_slide = first_new and i == 0 and len(s["points"]) <= 1
        slide = prs.slides.add_slide(prs.slide_layouts[0 if title_slide else 1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = _plain(s["title"])
        body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
        if body is not None:
            tf = body.text_frame
            tf.clear()
            for j, (level, text) in enumerate(s["points"]):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                para.text = _plain(text)
                para.level = level
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
    prs.save(str(path))
    return len(prs.slides)
