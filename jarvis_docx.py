"""Neatly formatted Word (.docx) files from simple Markdown, used by write_file for .docx names.

Supported: # to #### headings, - / * bullets (indent 2+ spaces to nest), numbered lines (kept as
plain text with their own numbers, so a list restarting at 1 never runs on from an earlier one),
**bold**, *italic*, `code` inline, | tables | (a |---| line under the first row makes it a header),
> quotes, ``` code blocks ```, a line of just --- for a page break, and LaTeX math ($inline$, or a
line of just $$display$$) as native Word equations: fractions, roots, powers/subscripts, \\left( \\right),
upright trig/log names, Greek letters and common symbols."""
import re

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

# $math$ must hug its dollars and not be followed by a digit, so "$5 and $10" stays plain text.
_INLINE = re.compile(r"(\$\$[^$]+\$\$|\$(?!\s)[^$\n]+?(?<!\s)\$(?!\d)|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
_DISPLAY = re.compile(r"\s*\$\$(.+)\$\$\s*")
_TOK = re.compile(r"\\[A-Za-z]+|\\.|\s+|\d+(?:\.\d+)?|.")
_FUNCS = {"sin", "cos", "tan", "sec", "csc", "cot", "arcsin", "arccos", "arctan", "sinh", "cosh", "tanh",
          "log", "ln", "exp", "lim", "max", "min"}
_SYMBOLS = {"theta": "θ", "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "Delta": "Δ", "phi": "φ",
            "pi": "π", "omega": "ω", "lambda": "λ", "mu": "μ", "sigma": "σ", "cdot": "·", "times": "×",
            "div": "÷", "pm": "±", "mp": "∓", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "ne": "≠",
            "neq": "≠", "approx": "≈", "implies": "⟹", "Rightarrow": "⇒", "iff": "⟺", "to": "→",
            "rightarrow": "→", "infty": "∞", "circ": "°", "in": "∈", "therefore": "∴", "ldots": "…",
            "cdots": "⋯", "dots": "…", "quad": " ", "qquad": "  ", ",": " ",
            ";": " ", ":": " ", "!": "", "{": "{", "}": "}", "%": "%", "|": "‖"}
_DELIMS = {".": "", "\\{": "{", "\\}": "}", "\\|": "‖", "\\langle": "⟨", "\\rangle": "⟩"}
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _base_styles(doc) -> None:
    """House style for a new document: Calibri 11, 1" margins, airy but compact spacing."""
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in ((1, 18), (2, 14), (3, 12), (4, 11)):
        h = doc.styles[f"Heading {level}"]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True


def _m(tag, *children, **attrs):
    el = OxmlElement(f"m:{tag}")
    for k, v in attrs.items():
        el.set(qn(f"m:{k}"), v)
    for c in children:
        el.append(c)
    return el


def _mr(text: str, plain: bool = False):
    r = _m("r")
    if plain:
        r.append(_m("rPr", _m("sty", val="p")))
    t = _m("t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


class _Latex:
    """Tiny LaTeX-math -> Word equation (OMML) parser for the notation an assistant writes in notes.
    Unknown commands degrade to their name as upright text instead of failing."""

    def __init__(self, src: str):
        self.toks = _TOK.findall(src)
        self.i = 0
        self.func = False

    def _peek(self):
        while self.i < len(self.toks) and self.toks[self.i].isspace():
            self.i += 1
        return self.toks[self.i] if self.i < len(self.toks) else None

    def _take(self):
        t = self._peek()
        self.i += 1
        return t

    def seq(self, stop=("}",)):
        out = []
        while (t := self._peek()) is not None and t not in stop and t != "\\right":
            out.extend(self.scripted())
        return out

    def group(self):
        if self._peek() == "{":
            self.i += 1
            out = self.seq()
            if self._peek() == "}":
                self.i += 1
            return out
        return self.atom() if self._peek() is not None else []

    def scripted(self):
        base = self.atom()
        func = self.func
        sub = sup = None
        while self._peek() in ("^", "_"):
            op = self._take()
            if op == "^" and self._peek() == "\\circ":  # 75^\circ -> 75°, not a raised ring
                self.i += 1
                base.append(_mr("°"))
                continue
            g = self.group()
            sup, sub = (g, sub) if op == "^" else (sup, g)
        if sup is not None and sub is not None:
            base = [_m("sSubSup", _m("e", *base), _m("sub", *sub), _m("sup", *sup))]
        elif sup is not None:
            base = [_m("sSup", _m("e", *base), _m("sup", *sup))]
        elif sub is not None:
            base = [_m("sSub", _m("e", *base), _m("sub", *sub))]
        nxt = self._peek()
        if func and nxt and nxt not in ("\\left", "\\right") and (nxt[0].isalnum() or nxt.startswith("\\")):
            base.append(_mr(" "))  # "sin x", not "sinx"
        return base

    def _delim(self):
        t = self._take() or ""
        return _DELIMS.get(t, t)

    def _raw_group(self) -> str:
        if self._peek() != "{":
            return self._take() or ""
        self.i += 1
        depth, text = 1, []
        while self.i < len(self.toks):
            t = self.toks[self.i]
            self.i += 1
            depth += (t == "{") - (t == "}")
            if depth == 0:
                break
            text.append(t)
        return "".join(text)

    def atom(self):
        self.func = False
        t = self._take()
        if t == "{":
            self.i -= 1
            return self.group()
        if t in ("\\frac", "\\dfrac", "\\tfrac"):
            num = self.group()
            return [_m("f", _m("num", *num), _m("den", *self.group()))]
        if t == "\\sqrt":
            deg = []
            if self._peek() == "[":
                self.i += 1
                deg = self.seq(stop=("]",))
                if self._peek() == "]":
                    self.i += 1
            pr = _m("radPr") if deg else _m("radPr", _m("degHide", val="1"))
            return [_m("rad", pr, _m("deg", *deg), _m("e", *self.group()))]
        if t == "\\left":
            beg = self._delim()
            body = self.seq()
            end = ""
            if self._peek() == "\\right":
                self.i += 1
                end = self._delim()
            return [_m("d", _m("dPr", _m("begChr", val=beg), _m("endChr", val=end)), _m("e", *body))]
        if t in ("\\text", "\\mathrm", "\\textrm", "\\operatorname"):
            return [_mr(self._raw_group(), plain=True)]
        if t.startswith("\\"):
            name = t[1:]
            if name in _FUNCS:
                self.func = True
                return [_mr(name, plain=True)]
            return [_mr(_SYMBOLS[name])] if name in _SYMBOLS else [_mr(name, plain=True)]
        return [_mr({"-": "−", "*": "·"}.get(t, t))]


def _omath(src: str):
    p = _Latex(src)
    out = []
    while p._peek() is not None:
        out.extend(p.seq())
        if p._peek() is not None:  # stray } or \right: skip it
            p.i += 1
    return _m("oMath", *out)


def _add_inline(par, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            par._p.append(_omath(part.strip("$")))
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_table(doc, rows: list[str]) -> None:
    header = len(rows) > 1 and bool(_TABLE_SEP.match(rows[1]))
    data = [_cells(r) for i, r in enumerate(rows) if not (header and i == 1)]
    ncols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=ncols)
    table.style = "Table Grid"
    for r, row in enumerate(data):
        for c in range(ncols):
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            _add_inline(par, row[c] if c < len(row) else "")
            if header and r == 0:
                for run in par.runs:
                    run.bold = True
                _shade(cell, "D9E2F3")
    doc.add_paragraph()  # breathing room after the table


def write(path, content: str, append: bool = False) -> None:
    from pathlib import Path

    p = Path(path)
    doc = docx.Document(str(p)) if append and p.exists() else docx.Document()
    if not (append and p.exists()):
        _base_styles(doc)
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip().startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.3)
            run = par.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _add_table(doc, rows)
            continue
        elif not line.strip():
            pass
        elif m := _DISPLAY.fullmatch(line):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.35 if line[:1].isspace() else 0)
            par._p.append(_m("oMathPara", _omath(m.group(1))))
        elif re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", line):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            _add_inline(doc.add_heading(level=len(m.group(1))), m.group(2).strip())
        elif m := re.match(r"^(\s*)[-*•]\s+(.*)", line):
            nested = len(m.group(1).expandtabs(4)) >= 2
            _add_inline(doc.add_paragraph(style="List Bullet 2" if nested else "List Bullet"), m.group(2))
        elif m := re.match(r"^\s*>\s?(.*)", line):
            par = doc.add_paragraph(style="Quote")
            _add_inline(par, m.group(1))
        elif m := re.match(r"^(\s*)(\d+[.)])\s+(.*)", line):
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(0.35 + (0.3 if len(m.group(1)) >= 2 else 0))
            pf.first_line_indent = Inches(-0.25)
            pf.space_after = Pt(3)
            _add_inline(par, f"{m.group(2)} {m.group(3)}")
        else:
            _add_inline(doc.add_paragraph(), line.strip())
        i += 1
    doc.save(str(p))
