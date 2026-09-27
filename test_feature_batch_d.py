"""Feature batch 2026-09-27, Phase D (FEATURES.md): .pptx via write_file, weekly improvement report.
Temp paths/DB only."""

import sqlite3
import threading
from datetime import datetime

import pytest

import jarvis_improvement_report as rep
import jarvis_pptx


MD = """# Quarterly review
Hero, September

# Results
- Revenue up 12%
  - Mostly **new** customers
- Churn flat
Notes: mention the churn dip in August

---
Plain point on an untitled slide
"""


def test_pptx_parse_and_write(tmp_path):
    slides = jarvis_pptx.parse(MD)
    assert [s["title"] for s in slides] == ["Quarterly review", "Results", ""]
    assert slides[1]["points"] == [(0, "Revenue up 12%"), (1, "Mostly **new** customers"), (0, "Churn flat")]
    assert slides[1]["notes"] == ["mention the churn dip in August"]
    out = tmp_path / "deck.pptx"
    assert jarvis_pptx.write(out, MD) == 3
    assert jarvis_pptx.write(out, "# Extra\n- one more", append=True) == 4
    from pptx import Presentation
    prs = Presentation(str(out))
    assert prs.slides[1].shapes.title.text == "Results"
    assert "Mostly new customers" in prs.slides[1].placeholders[1].text_frame.text
    assert "churn dip" in prs.slides[1].notes_slide.notes_text_frame.text


def test_write_file_makes_pptx(monkeypatch, tmp_path):
    monkeypatch.setenv("JARVIS_MEMORY_DB_PATH", str(tmp_path / "j.db"))
    import jarvis
    res = jarvis._write_file_tool(str(tmp_path / "talk.pptx"), MD, False)
    assert "PowerPoint deck (3 slides)" in res and (tmp_path / "talk.pptx").exists()


def test_improvement_report_is_read_only_and_useful(tmp_path):
    path = tmp_path / "r.db"
    c = sqlite3.connect(path)
    now = datetime.now().isoformat(timespec="seconds")
    c.execute("CREATE TABLE action_audit (id INTEGER PRIMARY KEY, timestamp TEXT, transcript TEXT, tool_name TEXT, "
              "tool_input TEXT, result TEXT)")
    for r in ("Tool failed: timeout", "Tool failed: timeout", "ok", "Couldn't reach it"):
        c.execute("INSERT INTO action_audit (timestamp, transcript, tool_name, tool_input, result) VALUES (?, '', "
                  "'weather', '{}', ?)", (now, r))
    c.execute("CREATE TABLE dashboard_sessions (id INTEGER PRIMARY KEY, source TEXT, transcript TEXT, status TEXT, "
              "reply TEXT, started_at TEXT, ended_at TEXT)")
    for _ in range(3):
        c.execute("INSERT INTO dashboard_sessions (source, transcript, status, started_at) VALUES "
                  "('voice', 'Open my study playlist', 'done', ?)", (now,))
    c.execute("CREATE TABLE notification_stats (kind TEXT PRIMARY KEY, delivered INT, acted INT, dismissed INT, updated_at TEXT)")
    c.execute("INSERT INTO notification_stats VALUES ('network', 9, 1, 6, ?)", (now,))
    c.commit()
    before = c.execute("SELECT COUNT(*) FROM action_audit").fetchone()
    r = rep.build(lambda: sqlite3.connect(path), threading.Lock(), 7, [{"e2e_ms": 8000}, {"e2e_ms": 900}])
    text = rep.format_report(r)
    assert "weather failed 3 of 4" in text
    assert "open my study playlist" in text and "voice macro" in text
    assert "network announcements 6 times" in text and "1 recent voice replies" in text
    assert c.execute("SELECT COUNT(*) FROM action_audit").fetchone() == before
    empty = rep.build(lambda: sqlite3.connect(tmp_path / "e.db"), threading.Lock())
    assert rep.format_report(empty).startswith("Nothing to improve")
